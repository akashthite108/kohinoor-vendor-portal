from datetime import datetime
import glob
import os
import re
import shutil
import uuid
from flask import Flask, redirect, render_template, request, send_file, url_for
import openpyxl
import pandas as pd

app = Flask(__name__)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

SUBMISSIONS_DIR = os.path.join(BASE_DIR, 'submissions')
VENDOR_SLIPS_DIR = os.path.join(BASE_DIR, 'vendor_filled_slips')
COMPARISON_POOL_DIR = os.path.join(BASE_DIR, 'comparison_pool')
MASTER_DB_CSV = os.path.join(BASE_DIR, 'master_procurement_database.csv')
STAGING_LOG_PATH = os.path.join(BASE_DIR, 'master_staging_log.csv')

for directory in [SUBMISSIONS_DIR, VENDOR_SLIPS_DIR, COMPARISON_POOL_DIR]:
    os.makedirs(directory, exist_ok=True)

TEMPLATE_SHEET = 'Comparision Format Sheet.xlsx'
ACTIVE_RFQ_TEMPLATE = 'RFQ.xlx.xlsx'


def sanitize_filename(name):
    if not name:
        return "unknown"
    return re.sub(r'[^\w\-_.]', '_', str(name))


def robust_remove(filepath):
    import time
    for _ in range(5):
        try:
            if os.path.exists(filepath):
                os.remove(filepath)
            break
        except PermissionError:
            time.sleep(0.3)
        except Exception:
            break


def init_db():
    if not os.path.exists(STAGING_LOG_PATH):
        df = pd.DataFrame(columns=[
            'Timestamp', 'RFQ_Code', 'Vendor_Name', 'GSTIN', 'File_Path',
        ])
        df.to_csv(STAGING_LOG_PATH, index=False)
    rebuild_master_database()


def assign_category(desc):
    desc_lower = str(desc).lower()
    if 'block' in desc_lower or 'brick' in desc_lower or 'aac' in desc_lower:
        return 'Masonry & Blocks'
    elif 'adhesive' in desc_lower or 'tile' in desc_lower or 'mortar' in desc_lower:
        return 'Finishing & Adhesives'
    elif 'pin' in desc_lower or 'wedge' in desc_lower or 'mivan' in desc_lower or 'formwork' in desc_lower:
        return 'Formwork & Hardware'
    elif 'waterproof' in desc_lower:
        return 'Waterproofing'
    else:
        return 'General Civil Works'


def get_boq_rows(ws):
    """Dynamically scans and returns all BOQ item rows (13 to 17)."""
    boq_rows = []
    for r in range(13, 18):
        desc = ws.cell(row=r, column=2).value
        if desc and str(desc).strip() and not str(desc).startswith("TOTAL"):
            boq_rows.append(r)
    return boq_rows


def rebuild_master_database():
    slips = [f for f in glob.glob(os.path.join(VENDOR_SLIPS_DIR, '*.xlsx')) if os.path.basename(f).startswith('RFQ_') and not os.path.basename(f).startswith('~$')]
    pool_slips = [f for f in glob.glob(os.path.join(COMPARISON_POOL_DIR, '*.xlsx')) if os.path.basename(f).startswith('RFQ_') and not os.path.basename(f).startswith('Kohinoor_Comparative') and not os.path.basename(f).startswith('~$')]

    all_slips = list(set(slips + pool_slips))
    records = []
    seen_files = set()

    for f in all_slips:
        fname = os.path.basename(f)
        if fname in seen_files:
            continue
        seen_files.add(fname)
        wb = None
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            ws = wb.active

            vendor_name = ws['F4'].value
            rfq_code = ws['F5'].value
            submission_date = ws['F6'].value
            gstin = ws['F7'].value
            vendor_contact = ws['F8'].value
            vendor_phone = ws.cell(row=9, column=6).value or 'N/A'
            vendor_email = ws.cell(row=9, column=2).value or 'N/A'
            vendor_address = ws.cell(row=10, column=2).value or 'N/A'
            project = ws['B4'].value
            tower = ws['B5'].value

            transport_term = ws['E22'].value or 'No'
            gst_term = ws['E23'].value or '18% GST'
            delivery_timeline = ws['E32'].value or 'Within 7 Days'
            payment_terms = ws['E33'].value or '30 Days Credit'
            remarks = ws['E35'].value or 'N/A'

            for r in get_boq_rows(ws):
                desc = ws.cell(row=r, column=2).value
                if desc:
                    try:
                        qty = float(ws.cell(row=r, column=4).value or 0)
                    except:
                        qty = 0.0
                    unit = ws.cell(row=r, column=5).value
                    brand = ws.cell(row=r, column=6).value
                    try:
                        rate_r0 = float(ws.cell(row=r, column=7).value or 0)
                    except:
                        rate_r0 = 0.0
                    try:
                        rate_r1 = float(ws.cell(row=r, column=8).value or rate_r0)
                    except:
                        rate_r1 = rate_r0

                    amount = qty * rate_r1
                    category = assign_category(desc)

                    records.append({
                        'Submission_Date': submission_date,
                        'RFQ_Code': rfq_code,
                        'Project': project,
                        'Tower': tower,
                        'Vendor_Name': vendor_name,
                        'GSTIN': gstin,
                        'Vendor_Phone': vendor_phone,
                        'Vendor_Email': vendor_email,
                        'Vendor_Address': vendor_address,
                        'Remarks': remarks,
                        'Material_Category': category,
                        'Material_Description': desc,
                        'Quantity': float(qty),
                        'Unit': unit,
                        'Brand': brand,
                        'Rate_R0': float(rate_r0),
                        'Rate_R1': float(rate_r1),
                        'Basic_Amount': float(amount),
                        'Transport': transport_term,
                        'GST_Term': gst_term,
                        'Delivery': delivery_timeline,
                        'Payment_Terms': payment_terms,
                        'Source_File': fname,
                    })
        except Exception as e:
            print(f"Error reading slip {fname}: {e}")
        finally:
            if wb:
                wb.close()

    df = pd.DataFrame(records)
    if not df.empty:
        df['Quantity'] = pd.to_numeric(df['Quantity'], errors='coerce').fillna(0.0)
        df['Rate_R0'] = pd.to_numeric(df['Rate_R0'], errors='coerce').fillna(0.0)
        df['Rate_R1'] = pd.to_numeric(df['Rate_R1'], errors='coerce').fillna(0.0)
        df['Basic_Amount'] = pd.to_numeric(df['Basic_Amount'], errors='coerce').fillna(0.0)
    else:
        df = pd.DataFrame(columns=[
            'Submission_Date', 'RFQ_Code', 'Project', 'Tower', 'Vendor_Name', 'GSTIN',
            'Vendor_Phone', 'Vendor_Email', 'Vendor_Address', 'Remarks',
            'Material_Category', 'Material_Description', 'Quantity', 'Unit', 'Brand',
            'Rate_R0', 'Rate_R1', 'Basic_Amount', 'Transport', 'GST_Term', 'Delivery',
            'Payment_Terms', 'Source_File'
        ])
    df.to_csv(MASTER_DB_CSV, index=False)
    return df


@app.route('/', methods=['GET', 'POST'])
def vendor_portal():
    if request.method == 'POST':
        rfq_code = request.form.get('rfq_code') or request.form.get('rfq_code_hidden')
        if not rfq_code or rfq_code.strip() == '':
            rfq_code = f"RFQ-{uuid.uuid4().hex[:6].upper()}"
    else:
        rfq_code = f"RFQ-{uuid.uuid4().hex[:6].upper()}"

    if not os.path.exists(ACTIVE_RFQ_TEMPLATE):
        return f'Error: RFQ template {ACTIVE_RFQ_TEMPLATE} missing on server.', 500

    wb = openpyxl.load_workbook(ACTIVE_RFQ_TEMPLATE)
    ws = wb.active

    def safe_set_cell(row_num, col_num, val):
        cell = ws.cell(row=row_num, column=col_num)
        coord = cell.coordinate
        if isinstance(cell, openpyxl.cell.cell.MergedCell):
            for merged_range in ws.merged_cells.ranges:
                if coord in merged_range:
                    ws.cell(row=merged_range.min_row, column=merged_range.min_col, value=val)
                    return
        cell.value = val

    safe_set_cell(5, 6, rfq_code)

    if request.method == 'POST':
        safe_set_cell(4, 2, request.form.get('project_name', 'Kohinoor West View'))
        safe_set_cell(5, 2, request.form.get('tower', 'Tower A - Phase 2'))
        safe_set_cell(6, 2, request.form.get('enquiry_date', datetime.now().strftime('%Y-%m-%d')))
        safe_set_cell(7, 2, request.form.get('buyer_org', 'Kohinoor Group Pune'))
        safe_set_cell(8, 2, request.form.get('buyer_contact', 'Purchase Dept'))

        vendor_name = request.form.get('vendor_name', 'Unknown Vendor')
        vendor_contact = request.form.get('vendor_contact', '')
        vendor_phone = request.form.get('vendor_phone', '')
        vendor_email = request.form.get('vendor_email', '')
        vendor_address = request.form.get('vendor_address', '')
        comm_remarks = request.form.get('comm_remarks', 'N/A')
        submitted_rfq = rfq_code
        
        safe_set_cell(4, 6, vendor_name)
        safe_set_cell(5, 6, submitted_rfq)
        safe_set_cell(6, 6, request.form.get('quotation_date', datetime.now().strftime('%Y-%m-%d')))
        safe_set_cell(7, 6, request.form.get('gstin', ''))
        safe_set_cell(8, 6, vendor_contact)
        safe_set_cell(9, 6, vendor_phone)
        safe_set_cell(9, 2, vendor_email)
        safe_set_cell(10, 2, vendor_address)
        safe_set_cell(35, 5, comm_remarks)

        safe_set_cell(22, 5, request.form.get('comm_transport', 'No'))
        safe_set_cell(23, 5, request.form.get('gst_rate', '18% GST'))
        safe_set_cell(25, 5, request.form.get('comm_loading', 'Included'))
        safe_set_cell(26, 5, request.form.get('comm_unloading', 'Included'))
        safe_set_cell(27, 5, request.form.get('comm_insurance', 'Not Applicable'))
        safe_set_cell(28, 5, request.form.get('comm_erection', 'Included'))
        safe_set_cell(29, 5, request.form.get('comm_warranty', '1 Year Standard'))
        safe_set_cell(30, 5, request.form.get('comm_guarantee', '1 Year Standard'))
        safe_set_cell(31, 5, request.form.get('comm_supplied', 'Yes'))
        safe_set_cell(32, 5, request.form.get('comm_delivery', 'Within 7 Days'))
        safe_set_cell(33, 5, request.form.get('comm_payment', '30 Days Credit'))
        safe_set_cell(34, 5, request.form.get('comm_validity', '60 Days'))

        for clear_r in range(13, 18):
            safe_set_cell(clear_r, 2, None)
            safe_set_cell(clear_r, 4, None)
            safe_set_cell(clear_r, 5, None)
            safe_set_cell(clear_r, 6, None)
            safe_set_cell(clear_r, 7, None)
            safe_set_cell(clear_r, 8, None)

        descs = request.form.getlist('desc[]')
        qtys = request.form.getlist('qty[]')
        units = request.form.getlist('unit[]')
        brands = request.form.getlist('brand[]')
        rates = request.form.getlist('rate[]')

        for i in range(len(descs)):
            r_idx = 13 + i
            if r_idx > 17:
                break
            safe_set_cell(r_idx, 2, descs[i])
            try:
                safe_set_cell(r_idx, 4, float(qtys[i]))
            except:
                pass
            safe_set_cell(r_idx, 5, units[i])
            safe_set_cell(r_idx, 6, brands[i])
            try:
                rate_val = float(rates[i])
                safe_set_cell(r_idx, 7, rate_val)
                safe_set_cell(r_idx, 8, f"=D{r_idx}*G{r_idx}")
            except:
                pass

        timestamp_str = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        safe_vendor_name = sanitize_filename(vendor_name)
        safe_rfq_code = sanitize_filename(submitted_rfq)
        
        filename = f'RFQ_{safe_vendor_name}_{safe_rfq_code}_{timestamp_str}.xlsx'
        filepath = os.path.join(VENDOR_SLIPS_DIR, filename)
        wb.save(filepath)
        wb.close()

        rebuild_master_database()
        return render_template('success.html', rfq_code=submitted_rfq, vendor_name=vendor_name)

    wb.close()
    return render_template('index.html', rfq_code=rfq_code)


@app.route('/purchase-dashboard')
def purchase_dashboard():
    search_query = request.args.get('q', '').strip().lower()

    raw_slips = [f for f in glob.glob(os.path.join(VENDOR_SLIPS_DIR, '*.xlsx')) if not os.path.basename(f).startswith('~$')]
    raw_slips_info = []
    for f in raw_slips:
        wb = None
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            ws = wb.active
            v_name = str(ws['F4'].value or '')
            r_code = str(ws['F5'].value or '')
            
            if search_query and search_query not in v_name.lower() and search_query not in r_code.lower():
                continue

            raw_slips_info.append({
                'filename': os.path.basename(f),
                'vendor': v_name,
                'rfq': r_code,
                'date': ws['F6'].value,
            })
        except Exception as e:
            print(f"Error reading slip {os.path.basename(f)}: {e}")
        finally:
            if wb:
                wb.close()

    pool_files = glob.glob(os.path.join(COMPARISON_POOL_DIR, '*.xlsx'))
    pool_files = [f for f in pool_files if not os.path.basename(f).startswith('Kohinoor_Comparative') and not os.path.basename(f).startswith('~$')]
    pool_slips_info = [os.path.basename(f) for f in pool_files]

    comparison_matrix = []
    for f_path in pool_files:
        wb = None
        try:
            wb = openpyxl.load_workbook(f_path, data_only=True)
            ws = wb.active
            vendor_name = ws['F4'].value
            source_fname = os.path.basename(f_path)
            for r in get_boq_rows(ws):
                desc = ws.cell(row=r, column=2).value
                if desc and str(desc).strip():
                    qty = ws.cell(row=r, column=4).value or 0
                    unit = ws.cell(row=r, column=5).value
                    brand = ws.cell(row=r, column=6).value
                    rate_r0 = ws.cell(row=r, column=7).value or 0
                    rate_r1 = ws.cell(row=r, column=8).value or rate_r0
                    comparison_matrix.append({
                        'Source_File': source_fname,
                        'Vendor_Name': vendor_name,
                        'Material_Description': desc,
                        'Quantity': qty,
                        'Unit': unit,
                        'Brand': brand,
                        'Rate_R0': rate_r0,
                        'Rate_R1': rate_r1,
                    })
        except Exception as e:
            print(f"Error reading pool file {os.path.basename(f_path)}: {e}")
        finally:
            if wb:
                wb.close()

    df_master = rebuild_master_database()
    vendor_summary = []
    if not df_master.empty and 'GSTIN' in df_master.columns:
        if search_query:
            df_master = df_master[
                df_master['Vendor_Name'].str.lower().str.contains(search_query, na=False) |
                df_master['GSTIN'].str.lower().str.contains(search_query, na=False) |
                df_master['RFQ_Code'].str.lower().str.contains(search_query, na=False)
            ]
        grouped = df_master.groupby('GSTIN').agg(
            Vendor_Name=('Vendor_Name', 'first'),
            Vendor_Phone=('Vendor_Phone', 'first'),
            Vendor_Email=('Vendor_Email', 'first'),
            Vendor_Address=('Vendor_Address', 'first'),
            Total_Quotes=('RFQ_Code', 'nunique'),
            Total_Quoted_Value=('Basic_Amount', 'sum'),
            Materials_Supplied=('Material_Description', lambda x: ", ".join(x.unique())),
            Last_Submission=('Submission_Date', 'max')
        ).reset_index()
        vendor_summary = grouped.to_dict(orient='records')

    return render_template(
        'dashboard.html',
        raw_slips=raw_slips_info,
        pool_slips=pool_slips_info,
        comparison_matrix=comparison_matrix,
        vendor_summary=vendor_summary,
        search_query=search_query
    )


@app.route('/vendor-analytics/<gstin>')
def vendor_analytics(gstin):
    df = rebuild_master_database()
    if df.empty or 'GSTIN' not in df.columns:
        return "No procurement data available.", 404
    
    vendor_df = df[df['GSTIN'] == gstin]
    if vendor_df.empty:
        return "Vendor not found in database.", 404
        
    vendor_name = vendor_df['Vendor_Name'].iloc[0]
    total_spend = float(vendor_df['Basic_Amount'].sum())
    total_quotes = int(vendor_df['RFQ_Code'].nunique())
    records = vendor_df.to_dict(orient='records')
    
    chart_labels = vendor_df['Material_Description'].tolist()
    chart_rates = vendor_df['Rate_R1'].tolist()
    chart_qtys = vendor_df['Quantity'].tolist()
    chart_dates = vendor_df['Submission_Date'].tolist()

    return render_template(
        'vendor_analytics.html',
        gstin=gstin,
        vendor_name=vendor_name,
        total_spend=total_spend,
        total_quotes=total_quotes,
        records=records,
        chart_labels=chart_labels,
        chart_rates=chart_rates,
        chart_qtys=chart_qtys,
        chart_dates=chart_dates
    )


@app.route('/analytics')
def analytics():
    df = rebuild_master_database()
    if not df.empty:
        df['Basic_Amount'] = pd.to_numeric(df['Basic_Amount'], errors='coerce').fillna(0.0)
        df['Quantity'] = pd.to_numeric(df['Quantity'], errors='coerce').fillna(0.0)
        df['Rate_R0'] = pd.to_numeric(df['Rate_R0'], errors='coerce').fillna(0.0)
        df['Rate_R1'] = pd.to_numeric(df['Rate_R1'], errors='coerce').fillna(0.0)

    total_submissions = df['Source_File'].nunique() if not df.empty else 0
    total_rfqs = df['RFQ_Code'].nunique() if not df.empty else 0
    total_materials = df['Material_Description'].nunique() if not df.empty else 0
    total_quoted_value = f"{df['Basic_Amount'].sum():,.2f}" if not df.empty else '0.00'

    l1_data, variance_data, commercial_data, category_summary = [], [], [], []
    master_records = df.to_dict(orient='records') if not df.empty else []

    if not df.empty:
        df_materials = df[~df['Material_Description'].isin(['Commercial Parameter', 'Basic Value', 'nan', 'NaN', '']) & df['Material_Description'].notna()].copy()

        if not df_materials.empty:
            idx_min = df_materials.groupby('Material_Description')['Rate_R1'].idxmin()
            df_l1 = df_materials.loc[idx_min]
            l1_data = df_l1.to_dict(orient='records')

            df_var = df_materials.copy()
            df_var['L1_Baseline'] = df_var.groupby('Material_Description')['Rate_R1'].transform('min')
            df_var['Percentage_Difference_From_L1'] = ((df_var['Rate_R1'] - df_var['L1_Baseline']) / df_var['L1_Baseline'] * 100).round(2)
            variance_data = df_var.to_dict(orient='records')

        df_comm = df[['Vendor_Name', 'RFQ_Code', 'Transport', 'GST_Term', 'Delivery', 'Payment_Terms', 'Source_File']].drop_duplicates(subset=['Source_File'])
        commercial_data = df_comm.to_dict(orient='records')

        df_cat = df.groupby('Material_Category').agg(
            Total_Spend=('Basic_Amount', 'sum'),
            Items_Count=('Material_Description', 'count'),
            Avg_Rate=('Rate_R1', 'mean'),
            Min_Rate=('Rate_R1', 'min'),
            Max_Rate=('Rate_R1', 'max'),
        ).reset_index()
        category_summary = df_cat.to_dict(orient='records')

    return render_template(
        'analytics.html',
        total_submissions=total_submissions,
        total_rfqs=total_rfqs,
        total_materials=total_materials,
        total_quoted_value=total_quoted_value,
        l1_data=l1_data,
        variance_data=variance_data,
        commercial_data=commercial_data,
        category_summary=category_summary,
        master_records=master_records,
    )


@app.route('/move-to-pool', methods=['POST'])
def move_to_pool():
    filename = request.form.get('filename')
    src = os.path.join(VENDOR_SLIPS_DIR, filename)
    dst = os.path.join(COMPARISON_POOL_DIR, filename)
    if os.path.exists(src):
        if os.path.exists(dst):
            robust_remove(dst)
        try:
            shutil.move(src, dst)
        except PermissionError:
            shutil.copy(src, dst)
            robust_remove(src)
    rebuild_master_database()
    return redirect(url_for('purchase_dashboard'))


@app.route('/remove-from-pool', methods=['POST'])
def remove_from_pool():
    filename = request.form.get('filename')
    dst = os.path.join(COMPARISON_POOL_DIR, filename)
    robust_remove(dst)
    rebuild_master_database()
    return redirect(url_for('purchase_dashboard'))


@app.route('/update-negotiated-rate', methods=['POST'])
def update_negotiated_rate():
    source_file = request.form.get('source_file')
    material_desc = request.form.get('material_desc', '').strip()
    try:
        new_rate_r1 = float(request.form.get('rate_r1', 0))
    except:
        new_rate_r1 = 0

    file_path = os.path.join(COMPARISON_POOL_DIR, source_file)
    if not os.path.exists(file_path):
        file_path = os.path.join(VENDOR_SLIPS_DIR, source_file)

    if os.path.exists(file_path):
        wb = None
        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb.active
            for r in get_boq_rows(ws):
                cell_val = ws.cell(row=r, column=2).value
                if cell_val and str(cell_val).strip() == material_desc:
                    ws.cell(row=r, column=8, value=new_rate_r1)
                    break
            wb.save(file_path)
        except Exception as e:
            print(f"Error updating rate: {e}")
        finally:
            if wb:
                wb.close()

    rebuild_master_database()
    return redirect(url_for('purchase_dashboard'))


@app.route('/generate-comparison', methods=['POST'])
def generate_comparison():
    template_path = TEMPLATE_SHEET
    if not os.path.exists(template_path):
        return "<script>alert('Error: Comparison format template sheet missing.'); window.location.href='/purchase-dashboard';</script>", 500

    pool_files = glob.glob(os.path.join(COMPARISON_POOL_DIR, '*.xlsx'))
    pool_files = [f for f in pool_files if not os.path.basename(f).startswith('Kohinoor_Comparative') and not os.path.basename(f).startswith('~$')]
    
    if len(pool_files) < 2:
        return "<script>alert('Strict Rule Error: Comparison requires a minimum of 2 vendors (maximum 5). Please add more vendors to the comparison pool.'); window.location.href='/purchase-dashboard';</script>", 400
    if len(pool_files) > 5:
        pool_files = pool_files[:5]

    comparison_title = request.form.get('comparison_title', 'General Civil Materials Procurement Comparison')
    timestamp_code = datetime.now().strftime('%Y%m%d-%H%M%S')
    unique_comp_code = f'COMP-{timestamp_code}'

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    def safe_set(row_num, col_num, val):
        cell = ws.cell(row=row_num, column=col_num)
        coord = cell.coordinate
        if isinstance(cell, openpyxl.cell.cell.MergedCell):
            for merged_range in ws.merged_cells.ranges:
                if coord in merged_range:
                    ws.cell(row=merged_range.min_row, column=merged_range.min_col, value=val)
                    return
        cell.value = val

    project_name = ''
    tower_name = ''
    target_rfq_code = 'COMPARATIVE'

    for idx, file_path in enumerate(pool_files):
        # Vendor block columns: Vendor A starts at Col F (6), Vendor B at Col I (9), etc.
        start_col = 6 + (idx * 3)
        v_wb = None
        try:
            v_wb = openpyxl.load_workbook(file_path, data_only=True)
            v_ws = v_wb.active

            vendor_name = v_ws['F4'].value
            project_name = v_ws['B4'].value
            tower_name = v_ws['B5'].value
            target_rfq_code = v_ws['F5'].value or 'COMPARATIVE'

            # Set Vendor Name safely in Row 4
            safe_set(4, start_col, vendor_name)

            boq_rows = get_boq_rows(v_ws)
            for i, v_row in enumerate(boq_rows):
                target_row = 6 + i  # BOQ items in new template are rows 6 to 10
                if target_row > 10:
                    break

                desc = v_ws.cell(row=v_row, column=2).value
                qty = v_ws.cell(row=v_row, column=4).value
                unit = v_ws.cell(row=v_row, column=5).value
                brand = v_ws.cell(row=v_row, column=6).value
                rate_r0 = v_ws.cell(row=v_row, column=7).value or 0
                rate_r1 = v_ws.cell(row=v_row, column=8).value or rate_r0

                if desc and str(desc).strip():
                    safe_set(target_row, 2, desc)
                    safe_set(target_row, 3, brand)   # Make / Brand Offered in Column C
                    safe_set(target_row, 4, qty)     # Qty in Column D
                    safe_set(target_row, 5, unit)    # Unit in Column E
                    safe_set(target_row, start_col, rate_r0)      # Rate R0
                    safe_set(target_row, start_col + 1, rate_r1)  # Rate R1

            # Commercial terms mapping for new template rows
            comm_rows = {
                12: 'E22',  # Transport
                13: 'E23',  # GST
                15: 'E31',  # Earlier Supplied
                16: 'E25',  # Loading
                17: 'E27',  # Insurance
                18: 'E26',  # Unloading
                19: 'E28',  # Erection
                20: 'E29',  # Warranty
                21: 'E30',  # Guarantee
                23: 'E32',  # Delivery
                24: 'E33',  # Payment Terms
                25: 'E34',  # Price Validity
                26: 'E35',  # Remark
            }

            for comp_row, v_cell in comm_rows.items():
                term_val = v_ws[v_cell].value
                safe_set(comp_row, start_col, term_val)
        finally:
            if v_wb:
                v_wb.close()

    ws['B2'] = project_name
    ws['B3'] = tower_name
    ws['M2'] = datetime.now().strftime('%Y-%m-%d')
    ws['M3'] = 'Purchase Dept (Automated)'

    safe_target_rfq = sanitize_filename(target_rfq_code)
    version_num = 1
    while True:
        output_filename = f'Kohinoor_Comparative_{safe_target_rfq}_{unique_comp_code}_V{version_num}.xlsx'
        output_path = os.path.join(COMPARISON_POOL_DIR, output_filename)
        if not os.path.exists(output_path):
            break
        version_num += 1

    wb.save(output_path)
    wb.close()
    rebuild_master_database()
    return send_file(output_path, as_attachment=True)


ARCHIVE_DIR = os.path.join(BASE_DIR, 'archive')
os.makedirs(ARCHIVE_DIR, exist_ok=True)


@app.route('/archive-submission', methods=['POST'])
def archive_submission():
    filename = request.form.get('filename')
    src = os.path.join(VENDOR_SLIPS_DIR, filename)
    dst = os.path.join(ARCHIVE_DIR, filename)
    if os.path.exists(src):
        if os.path.exists(dst):
            robust_remove(dst)
        try:
            shutil.move(src, dst)
        except PermissionError:
            shutil.copy(src, dst)
            robust_remove(src)
    rebuild_master_database()
    return redirect(url_for('purchase_dashboard'))


if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 8080))
    app.run(host='0.0.0.0', port=port, debug=True)